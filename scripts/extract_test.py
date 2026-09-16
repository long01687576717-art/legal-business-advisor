import pdfplumber
import docx
from pathlib import Path

# 项目根目录（注意：你的文件夹叫 dos，不是 docs）
BASE_DIR = Path(r"D:\HuaweiMoveData\Users\重庆森林\Desktop\法商项目")
DOCS_DIR = BASE_DIR / "dos"


def extract_pdf(path):
    text_parts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                text_parts.append(t)
    return "\n".join(text_parts)


def extract_docx(path):
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if row_text:
                parts.append(row_text)
    return "\n".join(parts)


def extract(path):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path)
    elif suffix == ".docx":
        return extract_docx(path)
    else:
        raise ValueError(f"不支持：{suffix}")


def main():
    for sub in ["企业", "法律"]:
        folder = DOCS_DIR / sub
        if not folder.exists():
            print(f"文件夹不存在：{folder}")
            continue
        print(f"\n{'=' * 60}")
        print(f"扫描文件夹：{folder}")
        print(f"{'=' * 60}")
        for f in sorted(folder.iterdir()):
            if f.suffix.lower() not in [".pdf", ".docx"]:
                continue
            print(f"\n--- 文件：{f.name} ---")
            try:
                text = extract(f)
                print(f"提取字符数：{len(text)}")
                print(f"前 500 字预览：")
                print(text[:500])
            except Exception as e:
                print(f"提取失败：{e}")


if __name__ == "__main__":
    main()