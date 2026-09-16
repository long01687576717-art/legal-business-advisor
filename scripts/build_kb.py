# -*- coding: utf-8 -*-
"""
构建向量知识库：
读取 dos/ 下的文档 -> 切块 -> 向量化 -> 存入 Chroma
"""
import pdfplumber
import docx
from pathlib import Path
import chromadb
from sentence_transformers import SentenceTransformer

# ============ 路径配置 ============
BASE_DIR = Path(r"D:\HuaweiMoveData\Users\重庆森林\Desktop\法商项目")
DOCS_DIR = BASE_DIR / "dos"
KB_DIR = BASE_DIR / "kb"
CHROMA_DIR = KB_DIR / "chroma_db"

# ============ 切块参数 ============
CHUNK_SIZE = 500       # 每块约 500 字
CHUNK_OVERLAP = 50     # 相邻块重叠 50 字，避免上下文断裂

# ============ 嵌入模型 ============
EMBED_MODEL_NAME = "BAAI/bge-small-zh-v1.5"


# ---------- 文档提取 ----------
def extract_pdf(path):
    text_parts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                text_parts.append(t)
    return "\n".join(text_parts)


def extract_docx(path):
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if row_text:
                parts.append(row_text)
    return "\n".join(parts)


def extract(path):
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path)
    elif suffix == ".docx":
        return extract_docx(path)
    elif suffix == ".txt":
        return path.read_text(encoding="utf-8")
    else:
        raise ValueError(f"不支持：{suffix}")


# ---------- 切块 ----------
def split_text(text, chunk_size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    """简单按字符切块，块之间保留重叠。"""
    text = text.replace("\r\n", "\n").strip()
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(text):
            break
        start = end - overlap
    return chunks


# ---------- 主流程 ----------
def main():
    print("=" * 60)
    print("开始构建知识库")
    print("=" * 60)

    # 1. 加载嵌入模型
    print(f"\n[1/4] 加载嵌入模型：{EMBED_MODEL_NAME}")
    print("（首次运行会自动下载模型，约 100MB，请耐心等待）")
    model = SentenceTransformer(EMBED_MODEL_NAME)
    print("模型加载完成。")

    # 2. 初始化 Chroma
    CHROMA_DIR.mkdir(parents=True, exist_ok=True)
    client = chromadb.PersistentClient(path=str(CHROMA_DIR))

    # 删除旧 collection（如果存在），避免重复
    try:
        client.delete_collection("knowledge_base")
        print("\n[2/4] 已清空旧 collection。")
    except Exception:
        print("\n[2/4] 新建 collection。")

    collection = client.create_collection(
        name="knowledge_base",
        metadata={"hnsw:space": "cosine"},
    )

    # 3. 遍历文档，提取 + 切块
    print("\n[3/4] 提取文档并切块...")
    all_chunks = []
    all_metadatas = []
    all_ids = []

    for sub in ["企业", "法律"]:
        folder = DOCS_DIR / sub
        if not folder.exists():
            continue
        for f in sorted(folder.iterdir()):
            if f.suffix.lower() not in [".pdf", ".docx", ".txt"]:
                continue
            print(f"  处理：{f.name}")
            try:
                text = extract(f)
                chunks = split_text(text)
                print(f"    提取 {len(text)} 字符，切成 {len(chunks)} 块")
                for i, chunk in enumerate(chunks):
                    all_chunks.append(chunk)
                    all_metadatas.append({
                        "source": f.name,
                        "category": sub,
                        "chunk_index": i,
                    })
                    all_ids.append(f"{sub}_{f.stem}_{i}")
            except Exception as e:
                print(f"    处理失败：{e}")

    print(f"\n  总共 {len(all_chunks)} 块待向量化。")

    # 4. 向量化并入库（分批，避免内存爆）
    print("\n[4/4] 向量化并写入 Chroma（可能需要几分钟）...")
    batch_size = 64
    for i in range(0, len(all_chunks), batch_size):
        batch_chunks = all_chunks[i:i + batch_size]
        batch_metas = all_metadatas[i:i + batch_size]
        batch_ids = all_ids[i:i + batch_size]

        embeddings = model.encode(batch_chunks, show_progress_bar=False).tolist()
        collection.add(
            documents=batch_chunks,
            embeddings=embeddings,
            metadatas=batch_metas,
            ids=batch_ids,
        )
        print(f"  已完成 {min(i + batch_size, len(all_chunks))}/{len(all_chunks)}")

    print("\n" + "=" * 60)
    print(f"知识库构建完成！共 {len(all_chunks)} 块。")
    print(f"数据库位置：{CHROMA_DIR}")
    print("=" * 60)


if __name__ == "__main__":
    main()